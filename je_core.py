"""Transparent, deterministic journal-entry screening and selection logic.

The results are audit planning aids, not conclusions about fraud or assurance.
"""
from __future__ import annotations

import csv
import io
import math
import random
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

REQUIRED = ("journal_id", "posting_date", "account", "debit", "credit")
ALIASES = {
    "journal_id": ("journal id", "journal number", "je number", "je id", "entry no", "entry number", "journal_no", "journal"),
    "line_id": ("line number", "line no", "line id", "row id"),
    "posting_date": ("posting date", "post date", "date posted", "transaction date", "je date"),
    "account": ("account number", "account code", "gl account", "gl code", "account name"),
    "debit": ("debits", "dr", "debit amount", "dr amount"),
    "credit": ("credits", "cr", "credit amount", "cr amount"),
    "description": ("memo", "narrative", "explanation", "journal description"),
    "preparer": ("created by", "entered by", "user", "user id", "prepared by"),
    "entry_time": ("created time", "time entered", "posted time", "posting time"),
    "source": ("entry source", "journal source", "entry type", "origin"),
    "approval_status": ("approval", "approved", "review status"),
}
WEIGHTS = {
    "Unbalanced entry": 35,
    "Weekend posting": 15,
    "After-hours entry": 15,
    "Large value": 25,
    "Round amount": 10,
    "Manual source": 15,
    "Near fiscal year-end": 15,
    "Missing preparer": 10,
    "Not approved": 20,
    "Potential duplicate line": 20,
}


def _key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _canon_map(headers):
    lookup = {}
    for name in ALIASES:
        for alias in (name, *ALIASES[name]):
            lookup.setdefault(_key(alias), name)
    mapping = {}
    for header in headers:
        normalized = lookup.get(_key(header))
        if normalized and normalized not in mapping:
            mapping[normalized] = header
    missing = [field for field in REQUIRED if field not in mapping]
    if missing:
        raise ValueError("Required columns not found: " + ", ".join(missing) + ". See docs/DATA_DICTIONARY.md.")
    return mapping


def _number(value, row_num, name):
    if value is None or str(value).strip() == "":
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        val = Decimal(str(value))
    else:
        v = str(value).strip().replace(",", "").replace("$", "")
        if v.startswith("(") and v.endswith(")"):
            v = "-" + v[1:-1]
        try:
            val = Decimal(v)
        except InvalidOperation as exc:
            raise ValueError(f"Row {row_num}: invalid {name} amount ({value!r}).") from exc
    if not val.is_finite():
        raise ValueError(f"Row {row_num}: non-finite {name} amount.")
    return val


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%d-%b-%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _after_hours(value):
    if not str(value or "").strip():
        return False
    text = str(value).strip()
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M:%S %p"):
        try:
            hour = datetime.strptime(text, fmt).hour
            return hour < 7 or hour >= 19
        except ValueError:
            pass
    return False


def _text(value):
    return "" if value is None else str(value).strip()


def import_data(filename, content):
    """Return a list of source record dictionaries from CSV or XLSX bytes."""
    if len(content) > 20 * 1024 * 1024:
        raise ValueError("File exceeds the 20 MB local prototype limit.")
    lower = filename.lower()
    if lower.endswith(".csv"):
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = content.decode("cp1252")
        reader = csv.DictReader(io.StringIO(decoded))
        if not reader.fieldnames:
            raise ValueError("The CSV has no header row.")
        headers = list(reader.fieldnames)
        rows = list(reader)
    elif lower.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("Could not read the XLSX workbook. Use a valid, non-macro Excel file.") from exc
        try:
            sheet = book.active
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, None)
            if first is None:
                raise ValueError("The first worksheet is empty.")
            headers = [str(v).strip() if v is not None else "" for v in first]
            rows = [dict(zip(headers, row)) for row in iterator if any(v is not None and str(v).strip() for v in row)]
        finally:
            book.close()
    else:
        raise ValueError("Only .csv and .xlsx files are supported.")
    mapping = _canon_map(headers)
    if not rows:
        raise ValueError("No data rows found.")
    if len(rows) > 100_000:
        raise ValueError("Prototype limit: 100,000 lines per upload.")
    result = []
    for idx, source in enumerate(rows, start=2):
        record = {name: source.get(orig) for name, orig in mapping.items()}
        record["source_row"] = idx
        result.append(record)
    return result


def analyze(raw_rows, config=None):
    """Run completeness indicators and a documented risk-based sample at journal level."""
    cfg = config or {}
    sample_size = max(0, min(1000, int(cfg.get("sample_size", 15))))
    risk_fraction = min(1, max(0, float(cfg.get("risk_fraction", 0.6))))
    seed = int(cfg.get("seed", 2026))
    year_end = _date(cfg.get("fiscal_year_end"))
    toler = Decimal("0.01")
    rows = []
    exceptions = []
    fingerprints = set()
    seen_line_ids = set()
    journal_groups = defaultdict(list)
    for i, raw in enumerate(raw_rows, start=1):
        row_num = raw.get("source_row", i + 1)
        je = _text(raw.get("journal_id"))
        if not je:
            exceptions.append({"row": row_num, "journal_id": "", "type": "Missing journal ID", "detail": "Journal number not populated"})
            je = f"[MISSING-ID-ROW-{row_num}]"
        posted = _date(raw.get("posting_date"))
        if not posted:
            exceptions.append({"row": row_num, "journal_id": je, "type": "Missing/invalid posting date", "detail": str(raw.get("posting_date") or "")})
        debit = _number(raw.get("debit"), row_num, "debit")
        credit = _number(raw.get("credit"), row_num, "credit")
        if debit < 0 or credit < 0:
            exceptions.append({"row": row_num, "journal_id": je, "type": "Negative debit/credit", "detail": "Review original data sign conventions"})
        if debit != 0 and credit != 0:
            exceptions.append({"row": row_num, "journal_id": je, "type": "Both debit and credit", "detail": "One line has both debit and credit populated"})
        account = _text(raw.get("account"))
        if not account:
            exceptions.append({"row": row_num, "journal_id": je, "type": "Missing account", "detail": "GL account not populated"})
        lineid = _text(raw.get("line_id"))
        if lineid:
            if (je, lineid) in seen_line_ids:
                exceptions.append({"row": row_num, "journal_id": je, "type": "Duplicate line ID", "detail": f"Line ID {lineid}"})
            seen_line_ids.add((je, lineid))
        fingerprint = (je, str(posted), account, debit, credit, _text(raw.get("description")))
        duplicate = fingerprint in fingerprints
        if duplicate:
            exceptions.append({"row": row_num, "journal_id": je, "type": "Potential duplicate line", "detail": "Same JE/date/account/amount/description; inspect supporting records"})
        fingerprints.add(fingerprint)
        normalized = {
            "source_row": row_num, "journal_id": je, "line_id": lineid,
            "posting_date": posted.isoformat() if posted else "", "account": account,
            "debit": float(debit), "credit": float(credit),
            "description": _text(raw.get("description")), "preparer": _text(raw.get("preparer")),
            "entry_time": _text(raw.get("entry_time")), "source": _text(raw.get("source")),
            "approval_status": _text(raw.get("approval_status")),
            "duplicate": duplicate,
        }
        rows.append(normalized)
        journal_groups[je].append(normalized)
    magnitudes = []
    for grouped in journal_groups.values():
        magnitudes.append(max(sum(Decimal(str(r["debit"])) for r in grouped), sum(Decimal(str(r["credit"])) for r in grouped)))
    sorted_vals = sorted(magnitudes)
    large_threshold = sorted_vals[max(0, math.ceil(len(sorted_vals) * 0.90) - 1)] if sorted_vals else Decimal("0")
    entries = []
    for je, grouped in journal_groups.items():
        debit = sum(Decimal(str(r["debit"])) for r in grouped)
        credit = sum(Decimal(str(r["credit"])) for r in grouped)
        magnitude = max(debit, credit)
        dates = [date.fromisoformat(r["posting_date"]) for r in grouped if r["posting_date"]]
        flags = []
        if abs(debit - credit) > toler:
            flags.append("Unbalanced entry")
            exceptions.append({"row": "", "journal_id": je, "type": "Unbalanced entry", "detail": f"Dr {debit:.2f} vs Cr {credit:.2f}"})
        if any(d.weekday() >= 5 for d in dates):
            flags.append("Weekend posting")
        if any(_after_hours(r["entry_time"]) for r in grouped):
            flags.append("After-hours entry")
        if len(sorted_vals) >= 10 and magnitude >= large_threshold and magnitude > 0:
            flags.append("Large value")
        if magnitude >= 1000 and magnitude % 1000 == 0:
            flags.append("Round amount")
        if any(r["source"].casefold() in ("manual", "manual je", "spreadsheet", "excel") for r in grouped):
            flags.append("Manual source")
        if year_end and any(year_end - timedelta(days=4) <= d <= year_end for d in dates):
            flags.append("Near fiscal year-end")
        if all(not r["preparer"] for r in grouped):
            flags.append("Missing preparer")
        if any(r["approval_status"].casefold() in ("not approved", "rejected", "pending", "unapproved") for r in grouped):
            flags.append("Not approved")
        if any(r["duplicate"] for r in grouped):
            flags.append("Potential duplicate line")
        score = min(100, sum(WEIGHTS[f] for f in flags))
        entries.append({"journal_id": je, "posting_date": min(dates).isoformat() if dates else "",
                        "lines": len(grouped), "debit": float(debit), "credit": float(credit),
                        "magnitude": float(magnitude), "difference": float(debit - credit),
                        "risk_score": score, "risk_level": "High" if score >= 50 else ("Medium" if score >= 20 else "Low"),
                        "flags": flags, "selected_by": ""})
    entries.sort(key=lambda e: (-e["risk_score"], -e["magnitude"], e["journal_id"]))
    n = min(sample_size, len(entries))
    n_risk = min(n, math.ceil(n * risk_fraction))
    selected = [e["journal_id"] for e in entries[:n_risk]]
    available = [e["journal_id"] for e in entries if e["journal_id"] not in selected]
    rng = random.Random(seed)
    rest = rng.sample(available, min(n - len(selected), len(available)))
    selected_set = set(selected + rest)
    for entry in entries:
        if entry["journal_id"] in selected_set:
            entry["selected_by"] = "Risk-based" if entry["journal_id"] in selected else "Random"
    selected_entries = [e for e in entries if e["journal_id"] in selected_set]
    total_debit = sum(Decimal(str(row["debit"])) for row in rows)
    total_credit = sum(Decimal(str(row["credit"])) for row in rows)
    control_checks = []
    for label, actual, val in (("Line count", len(rows), cfg.get("expected_lines")),
                               ("Total debits", total_debit, cfg.get("expected_debits")),
                               ("Total credits", total_credit, cfg.get("expected_credits"))):
        if val is None or str(val).strip() == "":
            continue
        try:
            expected = Decimal(str(val).replace(",", ""))
        except InvalidOperation as exc:
            raise ValueError(f"Invalid expected control total: {label}.") from exc
        difference = Decimal(str(actual)) - expected
        control_checks.append({"check": label, "expected": float(expected), "actual": float(actual),
                               "difference": float(difference), "pass": abs(difference) <= (Decimal("0") if label == "Line count" else toler)})
    summary = {"lines": len(rows), "journal_entries": len(entries), "total_debits": float(total_debit),
               "total_credits": float(total_credit), "exceptions": len(exceptions),
               "selected": len(selected_entries), "high_risk": sum(e["risk_level"] == "High" for e in entries),
               "selected_amount": round(sum(e["magnitude"] for e in selected_entries), 2),
               "large_value_threshold": float(large_threshold) if len(entries) >= 10 else None,
               "source_completeness_not_proven": True}
    return {"summary": summary, "entries": entries, "lines": rows, "exceptions": exceptions,
            "samples": selected_entries, "control_checks": control_checks,
            "config": {"sample_size": sample_size, "risk_fraction": risk_fraction, "seed": seed,
                       "fiscal_year_end": year_end.isoformat() if year_end else "",
                       "expected_lines": cfg.get("expected_lines", ""),
                       "expected_debits": cfg.get("expected_debits", ""),
                       "expected_credits": cfg.get("expected_credits", "")}}
