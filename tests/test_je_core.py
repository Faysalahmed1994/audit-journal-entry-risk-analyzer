import csv
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from je_core import analyze, import_data
from report import export_xlsx, safe_cell
from openpyxl import load_workbook


def one(journal_id, account, debit=0, credit=0, **kwargs):
    return dict(journal_id=journal_id, posting_date="2026-12-31", account=account,
                debit=debit, credit=credit, preparer="audit-demo", source="System", **kwargs)


class AuditJETests(unittest.TestCase):
    def test_balanced_entry_not_unbalanced(self):
        result = analyze([one("100", "Cash", 120, 0), one("100", "Revenue", 0, 120)], {"sample_size": 1})
        self.assertEqual(result["summary"]["journal_entries"], 1)
        self.assertEqual(result["summary"]["total_debits"], 120)
        self.assertNotIn("Unbalanced entry", result["entries"][0]["flags"])

    def test_unbalanced_je_flag_and_exception(self):
        result = analyze([one("101", "Cash", 500, 0), one("101", "Revenue", 0, 450)])
        self.assertIn("Unbalanced entry", result["entries"][0]["flags"])
        self.assertTrue(any(x["type"] == "Unbalanced entry" for x in result["exceptions"]))

    def test_missing_id_not_merged(self):
        result = analyze([one("", "A", 20), one("", "B", credit=20)])
        self.assertEqual(result["summary"]["journal_entries"], 2)

    def test_reproducible_random_selection(self):
        rows = []
        for i in range(100):
            rows.extend([one(str(i), "Cash", 100+i), one(str(i), "Revenue", credit=100+i)])
        config = {"sample_size": 12, "risk_fraction": 0.3, "seed": 543}
        first = analyze(rows, config)
        second = analyze(rows, config)
        self.assertEqual([e["journal_id"] for e in first["samples"]], [e["journal_id"] for e in second["samples"]])
        self.assertEqual(first["summary"]["selected"], 12)
        self.assertEqual(len(set(e["journal_id"] for e in first["samples"])), 12)

    def test_csv_alias_mapping(self):
        data = b'JE Number,Post Date,Account Number,DR,CR\nA1,2026-11-18,1000,250,0\nA1,2026-11-18,4000,0,250\n'
        rows = import_data("journal.csv", data)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["journal_id"], "A1")

    def test_control_totals(self):
        result = analyze([one("X1", "Cash", 150), one("X1", "Sales", credit=150)],
                         {"expected_lines": "2", "expected_debits": "150", "expected_credits": "150"})
        self.assertTrue(all(x["pass"] for x in result["control_checks"]))

    def test_xlsx_export_has_workpaper_tabs(self):
        result = analyze([one("Z1", "Cash", 200), one("Z1", "Income", credit=200)], {"sample_size": 1})
        payload = export_xlsx(result)
        workbook = load_workbook(io.BytesIO(payload), read_only=True)
        self.assertIn("JE Sample Workpaper", workbook.sheetnames)
        self.assertIn("Control Totals", workbook.sheetnames)
        self.assertEqual(workbook["JE Sample Workpaper"]["A2"].value, "Z1")
        workbook.close()

    def test_spreadsheet_injection_prevention(self):
        self.assertEqual(safe_cell("=HYPERLINK('bad')"), "'=HYPERLINK('bad')")
        self.assertEqual(safe_cell("+1+1"), "'+1+1")


class LocalServerSecurityTests(unittest.TestCase):
    def test_rejects_foreign_host_and_non_json_posts(self):
        import threading, urllib.request, urllib.error
        from http.server import ThreadingHTTPServer
        from app import Handler
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        port = server.server_address[1]
        def call(headers, data=b"{}"):
            req = urllib.request.Request(f"http://127.0.0.1:{port}/api/export", data=data, headers=headers, method="POST")
            try:
                return urllib.request.urlopen(req).status
            except urllib.error.HTTPError as err:
                return err.code
        try:
            self.assertEqual(call({"Host": "evil.example", "Content-Type": "application/json"}), 403)
            self.assertEqual(call({"Content-Type": "text/plain"}), 403)
            self.assertEqual(call({"Content-Type": "application/json", "Origin": "https://evil.example"}), 403)
            self.assertEqual(call({"Content-Type": "application/json"}), 400)  # trusted, but nothing analyzed yet
        finally:
            server.shutdown()
            server.server_close()


if __name__ == "__main__":
    unittest.main()
