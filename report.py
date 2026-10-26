"""Generate a local, reviewable .xlsx audit workpaper from the screening output."""
from __future__ import annotations

import base64
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "16324F"
BLUE = "DDEBF6"


def safe_cell(value):
    """Prevent spreadsheet formula injection from untrusted client ledger text."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sheet(book, name, headers, records):
    ws = book.create_sheet(name)
    ws.append(headers)
    for record in records:
        ws.append([safe_cell(v) for v in record])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for idx, head in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = min(37, max(14, len(str(head)) + 3))
    return ws


def export_xlsx(result, evidence=None):
    book = Workbook()
    intro = book.active
    intro.title = "Overview"
    intro["A1"] = "AUDITJE | JOURNAL ENTRY TESTING WORKPAPER"
    intro["A1"].font = Font(size=17, bold=True, color=NAVY)
    intro.merge_cells("A1:D1")
    intro["A3"] = "Generated (local)"
    intro["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    intro["A4"] = "Purpose"
    intro["B4"] = "Preliminary journal-entry risk screening and sample selection. Not a fraud determination."
    intro["A5"] = "Completeness limitation"
    intro["B5"] = "Reconcile source GL population to independent system reports; import alone cannot establish completeness."
    r = 7
    for key, val in result["summary"].items():
        intro.cell(r, 1, key.replace("_", " ").title())
        intro.cell(r, 2, safe_cell(val))
        r += 1
    intro.column_dimensions["A"].width = 29
    intro.column_dimensions["B"].width = 100
    intro["B4"].alignment = Alignment(wrap_text=True)
    intro["B5"].alignment = Alignment(wrap_text=True)
    intro.row_dimensions[4].height = 29
    intro.row_dimensions[5].height = 40
    for row in [4, 5]:
        intro[f"A{row}"].font = Font(bold=True, color=NAVY)
    intro["A22"] = "Reviewer sign-off"
    intro["A23"] = "Prepared by:"
    intro["A24"] = "Date:"
    intro["A25"] = "Reviewed by:"
    intro["A26"] = "Review date:"
    fields = ["journal_id", "posting_date", "lines", "debit", "credit", "difference", "magnitude", "risk_score", "risk_level", "flags", "selected_by"]
    _sheet(book, "Entry Risk Register",
           ["JE ID", "Posting Date", "Lines", "Debits", "Credits", "Out-of-balance", "Magnitude", "Risk Score", "Risk Tier", "Risk Flags", "Sample Type"],
           [[item.get(f, "") if f != "flags" else "; ".join(item["flags"]) for f in fields] for item in result["entries"]])
    _sheet(book, "GL Lines", ["Source Row", "JE ID", "Line ID", "Posting Date", "Account", "Debit", "Credit", "Description", "Preparer", "Entry Time", "Source", "Approval"],
           [[item.get(field, "") for field in ("source_row", "journal_id", "line_id", "posting_date", "account", "debit", "credit", "description", "preparer", "entry_time", "source", "approval_status")] for item in result["lines"]])
    samples = _sheet(book, "JE Sample Workpaper",
           ["JE ID", "Date", "Magnitude", "Risk Score", "Risk Flags", "Selection Basis", "Voucher / Evidence Ref", "Voucher Amount", "Evidence Agrees?", "Authorization Verified?", "Audit Result", "Auditor Notes", "Prepared By", "Reviewed By"],
           [[item["journal_id"], item["posting_date"], item["magnitude"], item["risk_score"], "; ".join(item["flags"]), item["selected_by"], "", "", "", "", "", "", "", ""] for item in result["samples"]])
    for c in ("G", "H", "I", "J", "K", "L", "M", "N"):
        samples.column_dimensions[c].width = 22
    _sheet(book, "Data Quality", ["Source Row", "JE ID", "Indicator", "Review Detail"],
           [[ex["row"], ex["journal_id"], ex["type"], ex["detail"]] for ex in result["exceptions"]])
    _sheet(book, "Control Totals", ["Check", "Expected", "Imported", "Difference", "Agrees?"],
           [[c["check"], c["expected"], c["actual"], c["difference"], "Yes" if c["pass"] else "No"] for c in result["control_checks"]])
    _sheet(book, "Methodology", ["Rule / Setting", "Value", "Notes"],
           [["Risk: " + k, v, "Rule-based planning indicator; not a fraud conclusion"] for k, v in __import__("je_core").WEIGHTS.items()] +
           [["Sampling: " + k, str(v), "Risk-first followed by reproducible random selection"] for k, v in result["config"].items()] +
           [["Limitation", "Population completeness requires external reconciliation", "Evaluate source system, date coverage and excluded entries"]])
    evidence = evidence or []
    if evidence:
        ws = _sheet(book, "Evidence Images", ["JE ID", "Evidence Label", "Attachment"], [])
        anchor_row = 2
        for item in evidence[:30]:
            je = str(item.get("journal_id", ""))[:100]
            label = str(item.get("label", ""))[:100]
            uri = item.get("data_url", "")
            if not isinstance(uri, str) or not uri.startswith(("data:image/png;base64,", "data:image/jpeg;base64,")):
                continue
            try:
                binary = base64.b64decode(uri.split(",", 1)[1], validate=True)
                if len(binary) > 3 * 1024 * 1024:
                    continue
                img = XLImage(io.BytesIO(binary))
                factor = min(1.0, 520 / max(img.width, 1), 320 / max(img.height, 1))
                img.width *= factor
                img.height *= factor
                ws.cell(anchor_row, 1, safe_cell(je))
                ws.cell(anchor_row, 2, safe_cell(label))
                ws.add_image(img, f"C{anchor_row}")
                anchor_row += 20
            except Exception:
                continue
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 80
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
